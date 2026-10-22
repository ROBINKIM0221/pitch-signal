"""발표자료 v2 조립 테스트 (2026-10-22). 실행: python -m pytest -q tests/test_build_deck.py
기록지 판독 그림이 저장소 밖(C:/work/private)에 있어 그 파일이 없으면 건너뛴다."""
import re

import pytest

from tools import build_deck as bd


@pytest.mark.skipif(not bd.READ_GRID.exists(), reason="기록지 판독 그림(저장소 밖)이 없음")
def test_deck_has_ten_shown_slides_three_hidden_and_seven_minutes_of_notes(tmp_path):
    from pptx import Presentation
    prs = Presentation(bd.build(tmp_path / "d.pptx"))
    slides = list(prs.slides)
    hidden = [s._element.get("show") == "0" for s in slides]
    assert len(slides) == 13 and hidden == [False] * 10 + [True] * 3
    total = 0
    for s, h in zip(slides, hidden):
        note = s.notes_slide.notes_text_frame.text
        assert note.strip()
        m = re.match(r"\[(\d+):(\d+)", note)
        if not h:
            total += int(m.group(1)) * 60 + int(m.group(2))
    assert total == 7 * 60
    names = {r.font.name for s in slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert names <= {bd.TITLE, bd.BRUSH, bd.BODY}
