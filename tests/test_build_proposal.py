"""제안서 조립 도구 테스트 (템플릿 v3, 2026-10-22). 실행: python -m pytest -q tests/test_build_proposal.py"""
import importlib.util
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("build_proposal", ROOT / "tools" / "build_proposal.py")
bp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bp)

SAMPLE = """# 1. 문제

첫 문단 **굵게** 그리고 [링크](https://example.com) 끝. {tab:a}와 {fig:b}를 본다.
이어지는 줄.

- 항목 하나
- 항목 둘

1. 번호 하나

Table: tab:a | 표 제목
| 열 A | 열 B |
|---|---|
| 1 | 2 |
Note: 주: 표 주석

![fig:b|캡션 {tab:a}](reports/figures/final/08_hs_gap.png){width=12cm}

```
C_t = P_t + Σ_e	(1)
```

> 강조 문단

# 2. 둘째 {nobreak}

## 2.1 소제목

[1] 첫 출처. https://a.io
"""


def test_parse_blocks_recognises_every_block_kind_in_order():
    kinds = [b[0] for b in bp.parse_blocks(SAMPLE)]
    assert kinds == ["h1", "p", "ul", "ol", "table", "note", "img", "code", "quote", "h1", "h2", "refs"]


def test_parse_blocks_reads_captions_labels_and_nobreak():
    blocks = bp.parse_blocks(SAMPLE)
    by = {b[0]: b for b in blocks}
    assert by["p"][1].endswith("이어지는 줄.")
    assert by["table"][1:] == ([["열 A", "열 B"], ["1", "2"]], "tab:a", "표 제목")      # 구분선(---)은 행이 아니다
    assert by["img"][1:] == ("캡션 {tab:a}", "reports/figures/final/08_hs_gap.png", 12.0, "fig:b")
    assert by["note"][1] == "주: 표 주석" and by["quote"][1] == "강조 문단"
    assert [b[2] for b in blocks if b[0] == "h1"] == [False, True]


def test_labels_are_numbered_in_order_and_unknown_labels_stop_the_build():
    labels = bp.number_labels(bp.parse_blocks(SAMPLE))
    assert labels == {"tab:a": "표 1", "fig:b": "그림 1"}
    assert bp.resolve_refs("{tab:a}와 {fig:b}", labels) == "표 1와 그림 1"
    with pytest.raises(KeyError):
        bp.resolve_refs("{fig:none}", labels)


def test_inline_runs_split_bold_links_and_bare_urls():
    runs = bp.inline_runs("A **B** [C](https://c.io) https://d.io 끝")
    assert runs == [("A ", False, None), ("B", True, None), (" ", False, None), ("C", False, "https://c.io"),
                    (" ", False, None), ("https://d.io", False, "https://d.io"), (" 끝", False, None)]
    assert bp.inline_runs("[1] 출처 [실행 뒤 기입]") == [("[1] 출처 [실행 뒤 기입]", False, None)]
    assert bp.inline_runs("지목했다 [10][11](표 4).") == [("지목했다 [10][11](표 4).", False, None)]          # 참고문헌 뒤 괄호는 링크가 아니다


def test_inline_math_is_its_own_run_kind():
    assert bp.inline_runs("값 $m_t$와 $Σ_e$") == [("값 ", False, None), ("m_t", False, "math"), ("와 ", False, None), ("Σ_e", False, "math")]


def test_equation_runs_mark_variables_subscripts_and_words():
    assert bp.equation_runs("x̄_t − min(0, z_{t−1})^{2}") == [
        ("x̄", "i"), ("t", "sub"), (" − ", ""), ("min", ""), ("(0, ", ""), ("z", "i"), ("t−1", "sub"), (")", ""), ("2", "sup")]


def test_build_fills_the_template_with_styles_numbers_and_table_borders(tmp_path):
    md = tmp_path / "p.md"
    md.write_text(SAMPLE, encoding="utf-8")
    out = bp.build(md, tmp_path / "p.docx", root=ROOT)
    from docx import Document
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert any(t.startswith("피치시그널") for t in texts) and bp.COVER["subtitle"] in texts and bp.COVER["field"] in texts   # 템플릿 표지
    assert not any("서식 기준" in t for t in texts)                                       # 템플릿 예시 본문은 지움
    sid = lambda p: p._p.pPr.pStyle.val if p._p.pPr is not None and p._p.pPr.pStyle is not None else None
    styles = {p.text: sid(p) for p in doc.paragraphs}
    assert styles["1. 문제"] == "Heading1" and styles["2.1 소제목"] == "Heading2"
    assert "표 1. 표 제목" in texts and "그림 1. 캡션 표 1" in texts and "주: 표 주석" in texts
    assert any("표 1와 그림 1" in t for t in texts)
    assert styles["표 1. 표 제목"] == "TableCaption" and styles["그림 1. 캡션 표 1"] == "FigureCaption"
    assert len(doc.inline_shapes) == 1 and len(doc.tables) == 1 and doc.tables[0].cell(1, 1).text == "2"
    head = doc.tables[0].cell(0, 0)._tc.tcPr.find(bp.qn("w:tcBorders"))
    assert head.find(bp.qn("w:top")).get(bp.qn("w:sz")) == "8" and head.find(bp.qn("w:bottom")).get(bp.qn("w:sz")) == "4"
    h1 = [p for p in doc.paragraphs if sid(p) == "Heading1"]
    assert [p.paragraph_format.page_break_before for p in h1] == [True, None]           # {nobreak}면 쪽을 넘기지 않는다
    ref = next(p for p in doc.paragraphs if p.text.startswith("[1]"))
    assert sid(ref) == "Reference" and ref.text.startswith("[1]\t")
