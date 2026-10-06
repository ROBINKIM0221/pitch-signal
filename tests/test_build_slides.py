"""발표자료 조립 도구 테스트 (단계 12). 실행: python -m pytest -q tests/test_build_slides.py"""
import importlib.util
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("build_slides", ROOT / "tools" / "build_slides.py")
bs = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bs)

SAMPLE = """# 표지
time: 0:10
layout: cover
message: 피치시그널
sub: 김형준 · UNIST
notes: 안녕하세요.

---
# 문제
time: 0:35
big: 35.6%
message: 고교 투수 3명 중 1명
- 하루 단위 규정
- ACWR 계산 가능 구간 4.9%
image: reports/figures/final/07_kbo_timeline.png
source: 야구공작소 2026
notes: 첫 문장.
둘째 문장.
"""


def test_parse_slides_reads_fields_bullets_and_multiline_notes():
    slides = bs.parse_slides(SAMPLE)
    assert [s["title"] for s in slides] == ["표지", "문제"]
    assert slides[0]["layout"] == "cover" and slides[0]["time"] == "0:10"
    assert slides[1]["big"] == "35.6%" and slides[1]["bullets"] == ["하루 단위 규정", "ACWR 계산 가능 구간 4.9%"]
    assert slides[1]["image"] == "reports/figures/final/07_kbo_timeline.png" and slides[1]["source"] == "야구공작소 2026"
    assert slides[1]["notes"] == "첫 문장.\n둘째 문장."


def test_total_time_adds_the_slide_times():
    assert bs.total_seconds(bs.parse_slides(SAMPLE)) == 45


def test_build_writes_a_deck_with_one_slide_per_entry_and_notes(tmp_path):
    md = tmp_path / "s.md"; md.write_text(SAMPLE, encoding="utf-8")
    out = bs.build(md, tmp_path / "s.pptx", root=ROOT)
    from pptx import Presentation
    prs = Presentation(out)
    assert len(prs.slides) == 2
    assert prs.slides[1].notes_slide.notes_text_frame.text == "첫 문장.\n둘째 문장."
    texts = " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
    assert "35.6%" in texts and "하루 단위 규정" in texts and "야구공작소 2026" in texts
    assert any(sh.shape_type == 13 for sh in prs.slides[1].shapes)          # 13 = 그림


def test_parse_slides_turns_escaped_newlines_in_field_values_into_line_breaks():
    assert bs.parse_slides("# A\nsub: 첫 줄\\n둘째 줄\n")[0]["sub"] == "첫 줄\n둘째 줄"
