"""제안서 조립: docs/proposal/proposal.md를 사용자 템플릿(docs/proposal/template.docx, 2026-10-22 v3)에 채워 Word로 만든다.

템플릿의 표지·목차(필드)·머리글·쪽 번호·스타일(본문, 제목 1~3, Table Caption, Table Text, Note, Figure, Figure Caption, Reference)을
그대로 쓰고, 본문은 템플릿 '서식 기준'에 맞춘다: 강조는 굵게만, 표는 위아래 굵은 선과 머리행 아래 얇은 선만, 그림은 Figure 문단.
Markdown은 아래 부분집합만 쓴다:
  #/##/### 제목 (끝에 {nobreak}를 붙이면 제목 1 앞에서 쪽을 넘기지 않음), 문단, '- ' 글머리, '1. ' 번호,
  'Table: 이름 | 제목' 다음 줄부터 | 표 |, 표 바로 뒤 'Note: 주석', ![이름|캡션](경로){width=15cm},
  ``` 수식 ``` (줄마다 '식<TAB>(번호)', _x·_{..} 아래 첨자, ^{..} 위 첨자), '> ' 강조 문단, [n] 참고문헌, <<<pagebreak>>>.
본문 안에서는 **굵게**, [글](주소), http 주소를 링크로 바꾸고, {fig:이름}·{tab:이름}을 '그림 N'·'표 N'으로 바꾼다(등장 순서로 번호).

사용 예:
    python tools/build_proposal.py --out "C:/work/private/제출물/PitchSignal_김형준_제안서.docx" --pdf
"""
from __future__ import annotations

import argparse
import copy
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "proposal" / "proposal.md"
TEMPLATE = ROOT / "docs" / "proposal" / "template.docx"
COVER = {   # 템플릿 표지에서 바꾸는 줄 (나머지 — 행사명·제안자·연락처·링크·날짜 — 은 템플릿 그대로)
    "subtitle": "고교부터 프로까지, 투수의 혹사와 부상 위험 신호를 알려 주는 모니터링 서비스",
    "field": "공모 분야: 경기력 향상 기술(부상 예방) · 대회 운영·관리 기술(경기기록 판독 자동화, 선수 관리)",
}
TEXT_W = 9072                                   # 본문 폭(dxa) = A4 21cm − 좌우 여백 2.5cm×2
BULLET_NUM = "3"                                # 템플릿 numbering.xml: numId 3 = 글머리(•), abstractNum 3 = 번호(1.)
DECIMAL_ABSTRACT = "3"


# ---------- Markdown 읽기 ----------

def parse_blocks(text: str) -> list[tuple]:
    """Markdown 부분집합을 블록 목록으로. 문단 안의 줄바꿈은 띄어쓰기로 잇는다."""
    blocks: list[tuple] = []
    lines = text.splitlines()
    i = 0
    para: list[str] = []
    caption: tuple[str, str] | None = None       # 'Table:' 줄이 다음 표에 붙일 (이름, 제목)

    def flush():
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush(); i += 1; continue
        if s == "<<<pagebreak>>>":
            flush(); blocks.append(("pagebreak",)); i += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*?)\s*(\{nobreak\})?$", s)
        if m:
            flush(); blocks.append((f"h{len(m.group(1))}", m.group(2).strip(), bool(m.group(3)))); i += 1; continue
        if s.startswith("```"):
            flush(); i += 1; code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip()); i += 1
            blocks.append(("code", "\n".join(code).strip("\n"))); i += 1; continue
        m = re.match(r"^!\[(?:(fig:[\w-]+)\|)?(.*?)\]\((.*?)\)(?:\{width=([\d.]+)cm\})?$", s)
        if m:
            flush(); blocks.append(("img", m.group(2), m.group(3), float(m.group(4)) if m.group(4) else 16.0, m.group(1))); i += 1; continue
        m = re.match(r"^Table:\s*(tab:[\w-]+)\s*\|\s*(.*)$", s)
        if m:
            flush(); caption = (m.group(1), m.group(2).strip()); i += 1; continue
        if s.startswith("Note:"):
            flush(); blocks.append(("note", s[5:].strip())); i += 1; continue
        if s.startswith("|"):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows, *(caption or (None, None)))); caption = None; continue
        if s.startswith("> "):
            flush(); blocks.append(("quote", s[2:].strip())); i += 1; continue
        if re.match(r"^\[\d+\]\s", s):
            flush(); refs = []
            while i < len(lines) and re.match(r"^\[\d+\]\s", lines[i].strip()):
                refs.append(lines[i].strip()); i += 1
            blocks.append(("refs", refs)); continue
        if re.match(r"^- ", s) or re.match(r"^\d+\. ", s):
            flush(); kind = "ul" if s.startswith("- ") else "ol"; items = []
            while i < len(lines) and (re.match(r"^- ", lines[i].strip()) if kind == "ul" else re.match(r"^\d+\. ", lines[i].strip())):
                items.append(re.sub(r"^(- |\d+\. )", "", lines[i].strip())); i += 1
            blocks.append((kind, items)); continue
        para.append(line); i += 1
    flush()
    return blocks


def number_labels(blocks: list[tuple]) -> dict[str, str]:
    """그림·표 이름에 등장 순서대로 번호: {'fig:x': '그림 1', 'tab:y': '표 1'}."""
    out, n = {}, {"fig": 0, "tab": 0}
    for b in blocks:
        label = b[4] if b[0] == "img" else b[2] if b[0] == "table" else None
        if label:
            kind = label.split(":")[0]
            n[kind] += 1
            out[label] = f"{'그림' if kind == 'fig' else '표'} {n[kind]}"
    return out


def resolve_refs(text: str, labels: dict[str, str]) -> str:
    """{fig:x}·{tab:y}를 '그림 N'·'표 N'으로. 없는 이름이면 멈춘다(오타가 그대로 나가지 않게)."""
    def sub(m):
        if m.group(1) not in labels:
            raise KeyError(f"정의되지 않은 그림·표 이름: {m.group(1)}")
        return labels[m.group(1)]
    return re.sub(r"\{((?:fig|tab):[\w-]+)\}", sub, text)


INLINE = re.compile(r"(\*\*.+?\*\*|\$[^$]+\$|\[[^\]]+\]\(https?://[^)]+\)|https?://[^\s)\]]+)")   # 링크는 주소가 http일 때만 — "[10](표 4)" 같은 참고문헌 뒤 괄호는 글


def inline_runs(text: str) -> list[tuple[str, bool, str | None]]:
    """본문 한 줄을 (글, 굵게, 링크 주소) 조각으로 나눈다. $..$ 수식은 링크 자리에 "math"를 둔다."""
    runs = []
    for part in INLINE.split(text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", part)
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, None))
        elif part.startswith("$") and part.endswith("$") and len(part) > 2:
            runs.append((part[1:-1], False, "math"))
        elif link:
            runs.append((link.group(1), False, link.group(2)))
        elif re.fullmatch(r"https?://\S+", part):
            runs.append((part, False, part))
        else:
            runs.append((part, False, None))
    return runs


def equation_runs(expr: str) -> list[tuple[str, str]]:
    """수식 한 줄을(글, 모양) 조각으로: 모양은 'i'(기울임 변수 한 글자, x̄처럼 붙은 기호 포함) · 'sub' · 'sup' · ''(그대로, min 같은 낱말 포함)."""
    out = []
    pattern = r"_\{([^}]*)\}|_([A-Za-z0-9−+]+)|\^\{([^}]*)\}|([A-Za-z]{2,})|([A-Za-z][̀-ͯ]*)|([^A-Za-z_^]+)"
    for m in re.finditer(pattern, expr):
        if m.group(1) is not None or m.group(2) is not None:
            out.append((m.group(1) if m.group(1) is not None else m.group(2), "sub"))
        elif m.group(3) is not None:
            out.append((m.group(3), "sup"))
        elif m.group(4):
            out.append((m.group(4), ""))
        elif m.group(5):
            out.append((m.group(5), "i"))
        else:
            out.append((m.group(6), ""))
    return out


# ---------- Word 쓰기 ----------

def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _set_style(paragraph, style_id: str):
    """스타일을 id로 붙인다 (템플릿 스타일 이름이 'Heading 1'처럼 대문자라 python-docx 이름 조회가 못 찾음)."""
    paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = style_id


def _para(doc: Document, style_id: str | None = None):
    p = doc.add_paragraph()
    if style_id:
        _set_style(p, style_id)
    return p


def _hyperlink(paragraph, url: str, text: str):
    """누를 수 있는 링크. 서식 기준(색 글자·밑줄 없음)을 따라 모양은 본문과 같다."""
    r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(t); link.append(run); paragraph._p.append(link)


def _write_inline(paragraph, text: str, bold_all: bool = False):
    for chunk, bold, url in inline_runs(text):
        if url == "math":
            _math_runs(paragraph, chunk)
        elif url:
            _hyperlink(paragraph, url, chunk)
        else:
            run = paragraph.add_run(chunk)
            if bold or bold_all:
                run.bold = True


def _cover(doc: Document):
    """템플릿 표지의 부제를 바꾸고 그 아래에 공모 분야 한 줄을 넣는다."""
    paras = doc.paragraphs
    title = next(i for i, p in enumerate(paras) if p.text.startswith("피치시그널"))
    sub = paras[title + 1]
    for r in sub.runs[1:]:
        r._r.getparent().remove(r._r)
    sub.runs[0].text = COVER["subtitle"]
    field = copy.deepcopy(sub._p)
    sub._p.addnext(field)
    from docx.text.paragraph import Paragraph
    fp = Paragraph(field, sub._parent)
    fp.runs[0].text = COVER["field"]; fp.runs[0].font.size = Pt(10.5)
    fp.paragraph_format.space_before = Pt(10); fp.paragraph_format.line_spacing = Pt(16.5)
    spacer = fp._p.getnext()                    # 아래 빈 문단의 간격을 늘어난 만큼(부제 두 줄 + 분야 두 줄) 줄여 표지를 한 쪽에 둔다
    Paragraph(spacer, sub._parent).paragraph_format.space_after = Pt(140)


def _clear_after_toc(doc: Document):
    """목차(sdt) 뒤 본문을 모두 지운다 (마지막 sectPr은 남김)."""
    body = doc.element.body
    sdt = body.find(qn("w:sdt"))
    for el in list(sdt.itersiblings()):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def _new_decimal_list(doc: Document) -> str:
    """번호 목록마다 1부터 다시 세는 새 번호 정의."""
    numbering = doc.part.numbering_part.element
    num_id = str(max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1)
    num = _el("w:num", numId=num_id); num.append(_el("w:abstractNumId", val=DECIMAL_ABSTRACT))
    ov = _el("w:lvlOverride", ilvl="0"); ov.append(_el("w:startOverride", val="1")); num.append(ov)
    numbering.append(num)
    return num_id


def _list_item(doc: Document, text: str, num_id: str):
    p = _para(doc, "ListParagraph")
    num_pr = _el("w:numPr"); num_pr.append(_el("w:ilvl", val="0")); num_pr.append(_el("w:numId", val=num_id))
    p._p.get_or_add_pPr().append(num_pr)
    _write_inline(p, text)


def _widths(rows: list[list[str]], ncol: int, total: int = TEXT_W, min_w: int = 1000) -> list[int]:
    """열마다 가장 긴 칸의 글자 수에 비례해 폭(dxa)을 나눈다 (너무 좁은 열은 최소 폭)."""
    longest = [max((len(r[c]) for r in rows if c < len(r)), default=1) for c in range(ncol)]
    weights = [min(max(w, 4), 40) ** 0.75 for w in longest]
    raw = [total * w / sum(weights) for w in weights]
    fixed = [max(w, min_w) for w in raw]
    scale = total / sum(fixed)
    w = [int(x * scale) for x in fixed]
    w[-1] += total - sum(w)
    return w


def _border(tag: str, size: int | None) -> OxmlElement:
    return _el(tag, val="single", color="000000", sz=size) if size else _el(tag, val="none", color="auto", sz=0)


def _table(doc: Document, rows: list[list[str]]):
    """템플릿 표 모양: 위 굵은 선, 머리행 아래 얇은 선, 맨 아래 굵은 선, 세로선·배경색 없음."""
    ncol = max(len(r) for r in rows)
    widths = _widths(rows, ncol)
    tbl = _el("w:tbl")
    pr = _el("w:tblPr"); pr.append(_el("w:tblW", type="dxa", w=TEXT_W))
    borders = _el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_border(f"w:{side}", None))
    pr.append(borders); tbl.append(pr)
    grid = _el("w:tblGrid")
    for w in widths:
        grid.append(_el("w:gridCol", w=w))
    tbl.append(grid)
    last = len(rows) - 1
    for ri, row in enumerate(rows):
        tr = _el("w:tr"); tr_pr = _el("w:trPr"); tr_pr.append(_el("w:tblHeader") if ri == 0 else _el("w:cantSplit")); tr.append(tr_pr)
        for ci in range(ncol):
            tc = _el("w:tc"); tc_pr = _el("w:tcPr"); tc_pr.append(_el("w:tcW", type="dxa", w=widths[ci]))
            tb = _el("w:tcBorders")
            tb.append(_border("w:top", 8 if ri == 0 else None)); tb.append(_border("w:left", None))
            tb.append(_border("w:bottom", 4 if ri == 0 else 8 if ri == last else None)); tb.append(_border("w:right", None))
            tc_pr.append(tb)
            mar = _el("w:tcMar")
            for side, w in (("top", 70), ("left", 100), ("bottom", 70), ("right", 100)):
                mar.append(_el(f"w:{side}", type="dxa", w=w))
            tc_pr.append(mar); tc_pr.append(_el("w:vAlign", val="center")); tc.append(tc_pr)
            tc.append(_el("w:p")); tr.append(tc)
        tbl.append(tr)
    doc.element.body.find(qn("w:sectPr")).addprevious(tbl)
    from docx.table import Table
    table = Table(tbl, doc._body)
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            p = table.cell(ri, ci).paragraphs[0]
            _set_style(p, "TableText"); p.alignment = 0
            _write_inline(p, row[ci] if ci < len(row) else "", bold_all=(ri == 0))
            if len(rows) <= 12 and ri < last:
                p.paragraph_format.keep_with_next = True      # 짧은 표는 쪽을 넘겨 한 덩어리로 (한 줄만 다음 쪽으로 떨어지지 않게)
    return table


def _caption(doc: Document, style: str, label: str, text: str):
    p = _para(doc, style)
    p.add_run(f"{label}.").bold = True
    _write_inline(p, " " + text)
    return p


def _math_runs(paragraph, expr: str):
    """수식 조각을 글 조각으로: 변수 한 글자는 기울임, 아래·위 첨자."""
    for text, shape in equation_runs(expr):
        run = paragraph.add_run(text)
        if shape in ("i", "sub", "sup"):
            run.italic = shape == "i" or bool(re.fullmatch(r"[A-Za-z]+", text))
        if shape == "sub":
            run.font.subscript = True
        if shape == "sup":
            run.font.superscript = True


def _equation(doc: Document, line: str):
    expr, _, num = line.partition("\t")
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.space_before = Pt(6); pf.space_after = Pt(10); pf.keep_with_next = True
    tabs = _el("w:tabs"); tabs.append(_el("w:tab", val="center", pos=TEXT_W // 2)); tabs.append(_el("w:tab", val="right", pos=TEXT_W))
    p._p.get_or_add_pPr().append(tabs)
    p.add_run("\t")
    _math_runs(p, expr.strip())
    if num:
        p.add_run("\t" + num.strip())


def build(md_path: Path, out: Path, root: Path = ROOT, template: Path = TEMPLATE) -> Path:
    """Markdown 원고를 템플릿에 채워 Word 문서를 쓴다. 그림 경로는 root 기준(절대 경로도 됨)이다."""
    blocks = parse_blocks(Path(md_path).read_text(encoding="utf-8"))
    labels = number_labels(blocks)
    ref = lambda s: resolve_refs(s, labels)
    doc = Document(str(template))
    _cover(doc)
    _clear_after_toc(doc)
    first_h1 = True
    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2", "h3"):
            p = _para(doc, f"Heading{kind[1]}")
            p.add_run(block[1])
            if kind == "h1" and (first_h1 or not block[2]):
                p.paragraph_format.page_break_before = True
            first_h1 = first_h1 and kind != "h1"
        elif kind == "p":
            _write_inline(doc.add_paragraph(), ref(block[1]))
        elif kind == "ul":
            for item in block[1]:
                _list_item(doc, ref(item), BULLET_NUM)
        elif kind == "ol":
            num_id = _new_decimal_list(doc)
            for item in block[1]:
                _list_item(doc, ref(item), num_id)
        elif kind == "refs":
            for line in block[1]:
                p = _para(doc, "Reference")
                _write_inline(p, re.sub(r"^(\[\d+\])\s+", "\\1\t", line))
        elif kind == "table":
            if block[2]:
                _caption(doc, "TableCaption", labels[block[2]], ref(block[3])).paragraph_format.keep_with_next = True
            _table(doc, [[ref(c) for c in r] for r in block[1]])
        elif kind == "note":
            _write_inline(_para(doc, "Note"), ref(block[1]))
        elif kind == "img":
            p = _para(doc, "Figure")
            p.add_run().add_picture(str(root / block[2]), width=Cm(min(block[3], 16.0)))
            if block[4]:
                _caption(doc, "FigureCaption", labels[block[4]], ref(block[1]))
            else:
                _write_inline(_para(doc, "FigureCaption"), ref(block[1]))
        elif kind == "code":
            for line in block[1].splitlines():
                _equation(doc, line)
        elif kind == "quote":
            p = doc.add_paragraph()
            pf = p.paragraph_format; pf.left_indent = pf.right_indent = Cm(1.0); pf.space_before = Pt(6); pf.space_after = Pt(10)
            _write_inline(p, ref(block[1]), bold_all=True)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    out = Path(out); out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def finalize(docx_path: Path, pdf_path: Path | None) -> Path | None:
    """설치된 Word로 목차·필드를 갱신해 docx에 저장하고, 원하면 PDF로도 내보낸다."""
    d = Path(docx_path).resolve()
    export = f"$d.ExportAsFixedFormat('{Path(pdf_path).resolve()}', 17); " if pdf_path else ""
    script = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; $d = $w.Documents.Open('{d}'); "
              f"$d.Fields.Update() | Out-Null; foreach ($t in $d.TablesOfContents) {{ $t.Update() }}; $d.Save(); {export}"
              f"$d.Close(0); $w.Quit()")
    before = Path(pdf_path).stat().st_mtime if pdf_path and Path(pdf_path).exists() else None
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    if pdf_path:
        after = Path(pdf_path).stat().st_mtime if Path(pdf_path).exists() else None
        if after is None or (before is not None and after <= before):
            raise RuntimeError(f"PDF가 저장되지 않았습니다: {pdf_path} — 다른 프로그램(뷰어)에 열려 있으면 닫고 다시 실행하거나 --pdf-name으로 다른 이름을 지정하세요")
    return pdf_path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=str(SOURCE))
    ap.add_argument("--template", default=str(TEMPLATE))
    ap.add_argument("--out", required=True, help="만들 .docx 경로")
    ap.add_argument("--pdf", action="store_true", help="같은 이름의 PDF도 만든다 (Word 필요)")
    ap.add_argument("--pdf-name", default=None, help="PDF 파일 이름을 따로 지정 (기본: docx와 같은 이름)")
    ap.add_argument("--no-word", action="store_true", help="Word로 목차를 갱신하지 않는다 (Word가 없는 환경)")
    args = ap.parse_args()
    out = build(Path(args.source), Path(args.out), template=Path(args.template))
    print(f"저장: {out}")
    if not args.no_word:
        pdf = (out.with_name(args.pdf_name) if args.pdf_name else out.with_suffix(".pdf")) if args.pdf else None
        finalize(out, pdf)
        print("목차 갱신" + (f", 저장: {pdf}" if pdf else ""))


if __name__ == "__main__":
    main()
