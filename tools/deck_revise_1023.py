"""발표자료 수정 (2026-10-23, 사용자와 10장을 함께 본 뒤 한 번에 — 메모 C:/work/private/pitchsignal/발표_수정메모.md).
사용자가 직접 고친 파일을 그대로 열어 고친다 (build_deck로 다시 만들면 사용자 수정이 사라짐). 영상은 이 다음에 PowerPoint로 넣는다.
    python -m tools.deck_revise_1023 <입력.pptx> <출력.pptx>
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from tools import build_deck as bd

DECK, PRIVATE = bd.DECK, bd.PRIVATE
READ_DECK = PRIVATE / "1021_판독과정" / "33829_95067_16_deck.png"


def drop(slide, *ids):
    for sh in [s for s in slide.shapes if s.shape_id in ids]:
        sh._element.getparent().remove(sh._element)


def shape(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)


def set_text(sh, new: str):
    """첫 문단 첫 조각 서식을 살려 글만 바꾼다 (나머지 조각·문단은 지움)."""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    p0.runs[0].text = new


def title(slide, t):
    set_text(shape(slide, 5), t)


def card(slide, x, y, w, h, num, unit, label, num_color=bd.BLUE, label_size=13, num_size=34, num_w=None):
    """찢어진 검정 상자: 왼쪽에 붓글씨 숫자, 오른쪽에 설명."""
    bd.picture(slide, DECK / "box_stat.png", x, y, w, h)
    nw = num_w or w * 0.36
    bd.text(slide, x + 0.3, y + 0.1, nw, h - 0.2, [[(num, bd.BRUSH, num_size, num_color), (unit, bd.TITLE, num_size * 0.55, num_color)]], anchor=MSO_ANCHOR.MIDDLE)
    bd.text(slide, x + 0.3 + nw, y + 0.12, w - nw - 0.55, h - 0.24, [[bd.body(t, label_size, bd.WHITE, i == 0)] for i, t in enumerate(label.split("\n"))],
            anchor=MSO_ANCHOR.MIDDLE, line=1.15)


def qr_corner(slide):
    bd.picture(slide, DECK / "qr.png", 11.9, 5.6, 1.0, 1.0)
    bd.text(slide, 11.45, 6.65, 1.85, 0.5, [[bd.body("robinkim0221.github.io", 8.5, bd.MUTED)], [bd.body("/pitch-signal", 8.5, bd.MUTED)]], align=PP_ALIGN.CENTER, line=1.0)


def s2(s):
    t = shape(s, 6); t.width = bd.Inches(12.13)
    set_text(t, "KBSA 규정은 하루 최대 105구와, 던진 투구 수에 따른 의무 휴식일(46~60구 1일 · 61~75구 2일 · 76~90구 3일 · 91~105구 4일)만 계산합니다")
    t.text_frame.paragraphs[0].runs[0].font.size = bd.Pt(16)
    r = shape(s, 11).text_frame.paragraphs[0].runs
    r[0].text = "272"
    o = shape(s, 14); o.left, o.top, o.width, o.height = bd.Inches(11.85), bd.Inches(6.05), bd.Inches(1.0), bd.Inches(1.0)   # 긴 설명 줄과 안 겹치게 오른쪽 아래로
    drop(s, 13)
    bd.text(s, 0.6, 5.65, 12.1, 1.0, [[bd.body("누적 부하 기준(피치시그널): 7일 투구 120구 이상, 또는 3일 내 두 경기 합 70구 이상 — 2025 전국 고교 투구 분포의 상위 5% · 20%", 14, bd.INK)],
                                      [bd.body("자료: KBSA 기록실, 2026 전국 고교 공식 경기 991경기 (피치시그널 집계)", 12, bd.MUTED)]], line=1.3)


def s3(s):
    title(s, "누구를 위한 서비스인가")
    for sid, t in ((8, "내 제구 흐름 확인, 경기 복기와 약점 파악"), (11, "오늘 누가 몇 구까지 · 팀 투수 누적 부하 관리"),
                   (14, "공식 기록에 없는 볼·스트라이크 비율과 경기 운영 능력")):
        sh = shape(s, sid); set_text(sh, t); sh.width = bd.Inches(5.3); sh.height = bd.Inches(0.75)
    for sid in (7, 10, 13):
        shape(s, sid).width = bd.Inches(5.3)
    pic = shape(s, 15)
    ratio = pic.height / pic.width
    pic.left, pic.top, pic.width = bd.Inches(7.45), bd.Inches(1.75), bd.Inches(5.4)
    pic.height = int(pic.width * ratio)


def s4(s):
    t = shape(s, 6)
    drop(s, *range(7, 17))
    set_text(t, "공식 기록에는 투구 수만 있고, 볼·스트라이크 개수는 협회 손기록지에만 남아 있습니다")
    r = t.text_frame.paragraphs[0].runs[0]; r.font.size = bd.Pt(16)
    t.top, t.height, t.width = bd.Inches(1.55), bd.Inches(0.8), bd.Inches(12.13)
    p = t.text_frame.add_paragraph()
    bd._font(p.add_run(), bd.BODY, 16, bd.MUTED); p.runs[0].text = "판독 모델이 손기록지 3,906장을 모두 읽고, 헷갈리는 칸은 AI 에이전트가 다시 읽어 볼과 스트라이크를 셌습니다"
    img = bd.framed("read_grid_deck", READ_DECK, private=True)
    h = 4.95; w = h / bd.fit(img, 1.0)
    bd.picture(s, img, 0.35, 2.4, h=h)
    x = 0.35 + w + 0.2
    for i, (num, unit, label) in enumerate((("3,906", "장", "협회 손기록지 사진 (2025·2026 시즌)"), ("98.7", "%", "2026 검산 통과 (40구 이상 등판)"),
                                            ("1.3", "%p", "중계 영상과 비교한 볼 비율 오차"))):
        bd.stat(s, x, 2.55 + i * 1.6, 12.85 - x, 1.42, num, unit, label)


def video_slide(s):
    """글·캡처를 지우고 영상 자리만 비워 둔다(영상은 PowerPoint로 넣음). QR과 주소만 작게."""
    drop(s, *[sh.shape_id for sh in s.shapes if sh.shape_id >= 5])
    qr_corner(s)


def s6(s):
    drop(s, 6, 7, 8)
    t = shape(s, 9); t.top, t.height = bd.Inches(0.75), bd.Inches(0.9)
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = bd.Pt(20); r.font.color.rgb = bd.WHITE; r.font.bold = True
    x = shape(s, 5); x.left, x.top, x.width, x.height = bd.Inches(0.45), bd.Inches(0.6), bd.Inches(1.0), bd.Inches(1.0)
    img = bd.framed("pitch_data", DECK / "pitch_data.png")
    w = 12.4; h = bd.fit(img, w)
    bd.picture(s, img, (bd.SW - w) / 2, 1.75 + max(0, (5.6 - h) / 2), w)


def s7(s):
    title(s, "다른 선수가 아닌, 평소의 나와 비교")
    drop(s, *range(6, 15))
    img = bd.framed("false_alarm_case", DECK / "false_alarm_case.png")
    w = 7.35
    bd.picture(s, img, 0.35, 1.6, w)
    h = bd.fit(img, w)
    bd.text(s, 0.6, 1.6 + h + 0.02, 7.0, 0.45, [[bd.body("부상 없던 Juan Mejia(2026, 불펜): 1마일 하락 경고선 20번 → 피치시그널 0번", 14, bd.INK, True)]])
    X = 7.95
    bd.text(s, X, 1.65, 4.9, 0.5, [[("왜 나와 비교하나", bd.TITLE, 20, bd.INK)]])
    bd.text(s, X, 2.15, 1.3, 0.9, [[("7", bd.BRUSH, 44, bd.BLUE), ("%", bd.TITLE, 26, bd.BLUE)]], anchor=MSO_ANCHOR.MIDDLE)
    bd.text(s, X + 1.35, 2.2, 3.55, 0.85, [[bd.body("구속·체중·나이로 선수끼리 비교하면 팔꿈치 수술의 7%만 설명됩니다", 14, bd.INK)]], anchor=MSO_ANCHOR.MIDDLE)
    bd.text(s, X, 3.35, 4.9, 0.5, [[("경보는 드물어야 믿는다", bd.TITLE, 20, bd.INK)]])
    card(s, X - 0.1, 3.9, 5.0, 1.3, "6.7", "번", "1마일 하락 경고선 (FanGraphs 구속 분석 기준)\n부상 없는 투수도 100등판에 6.7번\n불펜 한 명당 한 시즌 4번", num_color=bd.RGBColor(0xE8, 0x5A, 0x4F), label_size=12, num_w=1.45)
    card(s, X - 0.1, 5.3, 5.0, 1.3, "1", "번", "피치시그널\n정상이면 100등판에 1번\n한 시즌 0.3~0.5번", label_size=12, num_w=1.45)
    bd.text(s, 0.6, 6.85, 12.2, 0.5, [[bd.body("1마일(1.6km/h) 하락 경고선: FanGraphs(Petti, 2013) — 1마일 이상 떨어진 투수의 팔 부상률 11% vs 4%. 여기서는 시즌 초 기준 구간(선발 첫 8등판, 불펜 첫 15등판) 평균에 적용 · 자료: MLB 2021~2026 공개 투구 기록", 9.5, bd.MUTED)]])


def s9(s):
    title(s, "결과를 보기 전에 채점 기준부터 공개했다")
    drop(s, *range(6, 14))
    steps = (("01", "평가 계획을 먼저 공개", "2026.10.6 GitHub에 고정 · 무엇을 어떻게 잴지 미리 정함"),
             ("02", "2024~2025 시즌으로 한 번만 채점", "결과를 보고 기준을 바꾸지 않음"),
             ("03", "2026 시즌은 봉인했다가 한 번 채점", "10.7, 결과를 모르는 상태에서 실행"))
    for i, (n, head, desc) in enumerate(steps):
        y = 1.95 + i * 1.5
        bd.circle(s, 0.6, y, 0.9, n, size=24)
        bd.text(s, 1.75, y + 0.02, 4.6, 0.45, [[bd.body(head, 17, bd.INK, True)]])
        bd.text(s, 1.75, y + 0.47, 4.6, 0.7, [[bd.body(desc, 13, bd.MUTED)]])
    cards = (("3", "배", "부상 직전 5등판에 경보가 울린 비율\n부상 투수 10.3% vs 건강한 투수 3.3%"),
             ("2.2", "배", "경보 뒤 30일 안 팔 부상 IL\n10.5% vs 경보 없을 때 4.8% (2026도 2.0배)"),
             ("2", "등판", "먼저 울림\n첫 경보가 IL 등재보다 앞선 등판 수(중앙값)"),
             ("0.85", "회", "100등판당 실제 헛경보\n설계(1회)보다 적음 · 2026 봉인 시즌 0.78회"))
    for i, (num, unit, label) in enumerate(cards):
        card(s, 6.5, 1.75 + i * 1.25, 6.35, 1.15, num, unit, label, label_size=12.5, num_size=32, num_w=1.9)
    bd.text(s, 0.6, 6.85, 12.2, 0.4, [[bd.body("MLB 2021~2026 투수-시즌 3,040개 · 팔 부상 사례 310건을 같은 시점의 건강한 투수 2명씩과 짝지어 비교", 11, bd.MUTED)]])


def s10(s):
    title(s, "PitchSignal은 이렇게 쓰입니다")
    drop(s, *range(6, 15), 19)
    rows = (("대회 운영", "전국체전 · 주말리그 · 고교 전국대회(황금사자기·청룡기·대통령배·봉황대기·이마트배)에서 운영본부와 각 팀이 같은 화면으로 규정·누적 부하 확인. 대학리그·중학·리틀야구 대회에도 적용 가능"),
            ("협회 · 연맹", "손기록지 자동 판독으로 볼·스트라이크 기록 추가 · 며칠 단위 누적 기준 논의의 근거"),
            ("고교 · 아마추어 지도자", "오늘 누가 몇 구까지 던질 수 있는지 · 팀 투수 누적 부하 관리"),
            ("고교 · 아마추어 선수", "경기 복기(이닝별 투구 수·볼 비율) · 제구 흐름과 약점 파악 · 진학·드래프트용 시즌 기록"),
            ("프로 구단", "아마추어·해외 스카우트(외국인 투수 영입 전 점검) · 분석·트레이닝(구단 추적 데이터로 구속·폼 경보) · 불펜 운영"))
    y = 1.7
    for name, desc in rows:
        h = 1.05 if len(desc) > 70 else 0.85
        bd.picture(s, DECK / "box_small.png", 0.55, y + 0.05, 2.35, 0.55)
        bd.text(s, 0.7, y + 0.1, 2.1, 0.45, [[bd.body(name, 13, bd.WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        bd.text(s, 3.1, y + 0.05, 4.95, h, [[bd.body(desc, 13, bd.INK)]], line=1.15)
        y += h + 0.12
    bd.text(s, 0.6, 6.9, 7.6, 0.4, [[bd.body("연동: 트랙맨 등 투구 추적 장비 · KBO 고교 측정 프로그램 데이터 · 협회 기록실 기록지 사진", 11, bd.MUTED)]])


def main(src: str, out: str) -> None:
    prs = Presentation(src)
    sl = list(prs.slides)
    s2(sl[1]); s3(sl[2]); s4(sl[3]); video_slide(sl[4]); s6(sl[5]); s7(sl[6]); video_slide(sl[7]); s9(sl[8]); s10(sl[9])
    for s in sl:                                          # 사용자: 노트(메모)는 모두 지움
        if s.has_notes_slide:
            s.notes_slide.notes_text_frame.text = ""
        for sh in s.shapes:                               # 한글을 단어 단위로 줄바꿈 (기본값은 글자 중간에서 끊김: "가/능")
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    p._p.get_or_add_pPr().set("eaLnBrk", "0")
                    for run in p.runs:                    # 한국어로 표시해야 PowerPoint가 한글 단어 단위 줄바꿈을 적용
                        run._r.get_or_add_rPr().set("lang", "ko-KR")
    prs.save(out)
    print("저장:", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
