"""발표자료 v2 (2026-10-22, 사용자 결정: 'Beyond Goals' 템플릿 느낌, 10장 압축안, 7분, 고교 먼저 → 프로).
템플릿 PDF는 통째 그림이라 요소를 떼어 쓰지 않고 tools/deck_assets.py가 그린 종이 바탕·찢어진 상자·붓 자국을 쓴다.
글꼴: 한글 제목 Black Han Sans, 영어·숫자 붓글씨 Permanent Marker, 본문 Noto Sans KR (앞의 둘은 현재 사용자에 설치, pptx에 글꼴 포함 저장).
대본은 슬라이드 노트에 둔다 (영상 녹화와 11/20 발표에 그대로 읽는다). 숨김 백업 3장은 시연이 막힐 때만.

    python -m tools.build_deck --out "C:/work/private/제출물/PitchSignal_발표.pptx" --pdf
"""
from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from tools import deck_assets as A

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "reports" / "figures" / "deck"
FRAMED = DECK / "framed"
SCREENS = ROOT / "docs" / "proposal" / "screens"
FIG = ROOT / "reports" / "figures" / "final"
PRIVATE = Path("C:/work/private/pitchsignal/시안")
READ_GRID = PRIVATE / "1021_판독과정" / "33829_95067_16_grid.png"          # 기록지 사진 조각이라 저장소 밖
SERVICE = "https://robinkim0221.github.io/pitch-signal/"
TITLE, BRUSH, BODY = "Black Han Sans", "Permanent Marker", "Noto Sans KR"
INK, BLUE, WHITE, MUTED, RED, SOFT = (RGBColor(0x11, 0x11, 0x11), RGBColor(0x4B, 0xA3, 0xE6), RGBColor(0xFF, 0xFF, 0xFF),
                                       RGBColor(0x55, 0x55, 0x55), RGBColor(0xC8, 0x35, 0x2E), RGBColor(0xD8, 0xD8, 0xD8))
SW, SH = 13.333, 7.5


# ---------- 그리기 도구 ----------

def _font(run, name: str, size: float, color: RGBColor, bold: bool = False):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    latin = run._r.get_or_add_rPr().find(qn("a:latin"))
    for tag in ("a:cs", "a:ea"):                          # 한글(동아시아) 글꼴도 같은 이름으로 — latin 바로 뒤 ea, cs 순서
        el = OxmlElement(tag); el.set("typeface", name); latin.addnext(el)


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, rot=0.0, line=None, after=0):
    """paras: [[(글, 글꼴, 크기, 색, 굵게), ...], ...] — 문단마다 조각 목록."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        p.space_after = Pt(after)
        for t, name, size, color, *bold in runs:
            _font(p.add_run(), name, size, color, bool(bold and bold[0]))
            p.runs[-1].text = t
    box.rotation = rot
    return box


def body(t, size=16, color=INK, bold=False):
    return (t, BODY, size, color, bold)


def picture(slide, path: Path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def circle(slide, x, y, d, label: str, fill=BLUE, size=26):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill; shp.line.fill.background(); shp.shadow.inherit = False
    tf = shp.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _font(p.add_run(), BRUSH, size, WHITE); p.runs[0].text = label
    return shp


def background(slide, dark=False):
    picture(slide, DECK / ("paper_dark.png" if dark else "paper_light.png"), 0, 0, SW, SH)


def header(slide, dark=False):
    c = SOFT if dark else MUTED
    text(slide, 0.6, 0.3, 5, 0.3, [[("PITCHSIGNAL  피치시그널", BODY, 10, c)]])
    text(slide, SW - 5.6, 0.3, 5, 0.3, [[("2026 경기스포츠산업 공모전", BODY, 10, c)]], align=PP_ALIGN.RIGHT)


def title(slide, t: str, color=INK, y=0.72, size=40):
    text(slide, 0.6, y, SW - 1.2, 0.9, [[(t, TITLE, size, color)]])


def stat(slide, x, y, w, h, number: str, unit: str, label: str, box="box_stat.png", num_size=40, label_size=14):
    """찢어진 검정 상자 위에 파란 붓글씨 숫자와 흰 설명."""
    picture(slide, DECK / box, x, y, w, h)
    text(slide, x + 0.35, y + 0.18, w - 0.7, h * 0.55, [[(number, BRUSH, num_size, BLUE), (unit, TITLE, num_size * 0.62, BLUE)]], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.35, y + h * 0.6, w - 0.7, h * 0.35, [[body(label, label_size, WHITE)]])


def framed(name: str, src: Path, angle=0.0, crop_h: int | None = None, private=False) -> Path:
    """화면 캡처에 찢어진 종이 테두리 (필요하면 위쪽만 잘라서)."""
    out_dir = PRIVATE / "deck" if private else FRAMED
    dst = out_dir / f"{name}.png"
    if crop_h:
        im = Image.open(src)
        tmp = out_dir / f"_{name}_crop.png"; out_dir.mkdir(parents=True, exist_ok=True)
        im.crop((0, 0, im.width, min(crop_h, im.height))).save(tmp)
        A.frame(tmp, dst, angle=angle)
        tmp.unlink()
        return dst
    return A.frame(src, dst, angle=angle)


def fit(path: Path, w: float) -> float:
    im = Image.open(path)
    return w * im.height / im.width


def qr(path: Path) -> Path:
    import qrcode
    qrcode.make(SERVICE, box_size=10, border=2).save(path)
    return path


def notes(slide, t: str):
    slide.notes_slide.notes_text_frame.text = t


# ---------- 슬라이드 ----------

def s1_cover(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s, dark=True)
    picture(s, DECK / "stripes.png", 0, 0, 0.36, SH)
    text(s, SW - 5.6, 0.3, 5, 0.3, [[("2026 경기스포츠산업 공모전", BODY, 11, SOFT)]], align=PP_ALIGN.RIGHT)
    text(s, 0.85, 1.35, 6.4, 1.2, [[("피치시그널", TITLE, 72, WHITE)]])
    text(s, 0.95, 2.55, 6.0, 1.0, [[("PITCHSIGNAL", BRUSH, 46, BLUE)]], rot=-3)
    text(s, 0.9, 3.85, 6.0, 1.2, [[body("고교부터 프로까지,", 20, WHITE, True)], [body("투수의 혹사와 부상 위험 신호를 알려 줍니다", 20, WHITE, True)]], line=1.2)
    text(s, 0.9, 6.15, 6.0, 0.8, [[body("김형준 · UNIST", 15, WHITE, True)], [body("robinkim0221.github.io/pitch-signal", 12, BLUE)]], line=1.2)
    picture(s, DECK / "box_paper.png", 7.15, 1.2, 5.5, 4.77).rotation = 2
    text(s, 7.6, 1.7, 4.7, 1.6, [[("1,791", BRUSH, 96, INK), ("구", TITLE, 64, INK)]], rot=2)
    text(s, 7.7, 3.45, 4.6, 0.5, [[body("2026년 한 고교 투수가 한 시즌에 던진 공", 17, INK)]], rot=2)
    text(s, 7.7, 4.15, 4.6, 0.8, [[("규정이 문제 삼은 날은 하루", TITLE, 28, RED)]], rot=2)
    picture(s, DECK / "brush_x.png", 11.6, 6.0, 1.25, 1.25)
    notes(s, "[0:20] 1,791구. 2026년 한 고교 투수가 한 시즌에 던진 공입니다. 규정이 문제 삼은 날은 그중 하루뿐이었습니다. "
             "투수의 혹사와 부상 위험을 알려 주는 피치시그널, 김형준입니다.")


def s2_problem(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "규정은 하루만 본다")
    text(s, 0.6, 1.65, 11, 0.5, [[body("KBSA 규정은 하루 105구와 투구 수별 휴식일만 계산합니다", 18, MUTED)]])
    for x, num, unit, label in ((0.6, "1", "건", "2026 전국 고교 규정 위반"), (6.85, "270", "명", "규정은 지켰지만 누적 부하 기준을 넘은 투수")):
        picture(s, DECK / "box_wide.png", x, 2.55, 5.9, 2.75)
        text(s, x + 0.5, 2.85, 5.0, 1.4, [[(num, BRUSH, 84, BLUE), (unit, TITLE, 54, BLUE)]], anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.5, 4.35, 5.0, 0.7, [[body(label, 18, WHITE, True)]])
    text(s, 0.6, 5.7, 12, 0.9, [[body("누적 부하 기준: 7일 투구 120구 이상, 또는 3일 내 두 경기 합 70구 이상", 14, INK)],
                                [body("자료: KBSA 기록실, 2026 전국 고교 공식 경기 988경기 (피치시그널 집계)", 12, MUTED)]], line=1.3)
    picture(s, DECK / "brush_o.png", 11.55, 0.55, 1.3, 1.3)
    notes(s, "[0:35] 고교 투구수 규정은 하루 105구와 휴식일만 계산합니다. 며칠, 몇 주에 걸쳐 쌓이는 부하는 보지 않습니다. "
             "2026년 전국 고교 공식 경기를 모두 모아 보니 규정 위반은 1건이었습니다. 그런데 규정을 다 지키고도 누적 부하 기준을 넘은 투수는 270명이었습니다. "
             "규정을 지키는 것과 혹사하지 않는 것은 다릅니다.")


def s3_users(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "고교 야구의 세 사람")
    rows = (("01", "고교 선수", "내 누적 부하와 제구 흐름을 한 장에서"), ("02", "고교 지도자", "오늘 누가, 몇 구까지 던질 수 있는지"),
            ("03", "프로 구단 아마추어 스카우트", "투구 수만으로는 안 보이던 제구와 성적"))
    for i, (n, head, desc) in enumerate(rows):
        y = 2.0 + i * 1.5
        circle(s, 0.65, y, 1.0, n)
        text(s, 1.95, y + 0.05, 4.9, 0.5, [[body(head, 21, INK, True)]])
        text(s, 1.95, y + 0.55, 4.9, 0.5, [[body(desc, 16, MUTED)]])
    img = assets["card"]; picture(s, img, 6.75, 1.55, 6.1)
    notes(s, "[0:25] 그래서 고교 야구의 세 사람을 위해 만들었습니다. 선수는 자기 부하와 제구를, 지도자는 오늘 누가 몇 구까지 던질 수 있는지를, "
             "프로 구단 스카우트는 투구 수만으로는 알 수 없던 제구를 봅니다.")


def s4_reading(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "기록에 없는 볼을 읽다")
    text(s, 0.6, 1.65, 11.5, 0.5, [[body("공식 기록에는 투구 수만 있습니다. 볼과 스트라이크는 협회 손기록지에 남아 있습니다", 18, MUTED)]])
    h = 5.1; w = h / fit(assets["read"], 1.0)                 # 아래가 잘리지 않게 높이에 맞춘다
    picture(s, assets["read"], 0.45, 2.15, h=h)
    x = 0.45 + w + 0.45
    for i, (num, unit, label) in enumerate((("3,906", "장", "협회 손기록지 사진 (2025·2026 시즌)"), ("98.7", "%", "2026 검산 통과 (40구 이상 등판)"),
                                            ("1.3", "%p", "중계 영상과 비교한 볼 비율 오차"))):
        stat(s, x, 2.3 + i * 1.62, SW - 0.6 - x, 1.45, num, unit, label)
    notes(s, "[0:50] 그런데 공식 기록에는 볼과 스트라이크가 없습니다. 답은 협회 손기록지 사진에 있었습니다. 칸을 찾고, 기록원이 그린 표시를 하나씩 떼어 모델이 판정합니다. "
             "헷갈리는 칸은 AI가 칸 전체를 다시 읽고, 그리지 않은 마지막 공은 공식 기록으로 채웁니다. 이렇게 두 시즌 3,906장을 읽었고, 읽은 공 수는 공식 투구 수와 맞춰 검산합니다. "
             "중계 영상과 비교한 볼 비율 오차는 1.3%p입니다.")


def s5_demo_hs(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "시연 ① 고교 현황판")
    steps = ("선수 이름으로 찾기", "규정 · 누적 부하 신호등", "시즌 카드: 볼 비율 · 성적", "투수 운용 계획판")
    for i, t in enumerate(steps):
        y = 1.95 + i * 0.85
        circle(s, 0.65, y, 0.6, str(i + 1), size=20)
        text(s, 1.45, y + 0.08, 3.6, 0.5, [[body(t, 17, INK, True)]])
    picture(s, assets["qr"], 0.65, 5.45, 1.35, 1.35)
    text(s, 2.15, 5.75, 3.0, 0.8, [[body("바로 써 보기", 14, INK, True)], [body("robinkim0221.github.io/pitch-signal", 11, MUTED)]], line=1.2)
    picture(s, assets["hs"], 5.2, 1.35, 7.75)
    notes(s, "[1:15 시연] (선수 검색) 실제 화면입니다. 2025년 유신고 투수를 이름으로 찾겠습니다. "
             "(목록) 규정은 한 번도 어기지 않았지만 누적 부하는 '높음'입니다. 20번 등판해 1,134구, 일주일에 최대 206구를 던졌습니다. "
             "(시즌 카드) 기록지에서 읽은 볼 비율은 34%입니다. 스카우트는 이닝, 평균자책점, 삼진과 볼넷 비율도 같이 봅니다. "
             "(그래프) 7일 합이 주의선, 높음선을 넘는 구간이 바로 보입니다. "
             "(계획판) 지도자용 계획판입니다. 날짜를 고르면 그날 누가 나갈 수 있는지 나오고, 오늘 80구를 던지면 다음 등판이 언제 가능한지도 바로 알려 줍니다.")


def s6_bridge(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s, dark=True); header(s, dark=True)
    picture(s, DECK / "brush_x.png", 0.9, 1.0, 1.6, 1.6)
    picture(s, DECK / "brush_o.png", 10.9, 4.9, 1.7, 1.7)
    text(s, 1.5, 2.35, SW - 3.0, 1.0, [[("혹사를 피해 프로에 가도", TITLE, 46, WHITE)]], align=PP_ALIGN.CENTER)
    text(s, 1.5, 3.3, SW - 3.0, 1.0, [[("부상 위험은 계속된다", TITLE, 46, BLUE)]], align=PP_ALIGN.CENTER)
    text(s, 1.5, 4.65, SW - 3.0, 0.9, [[body("프로 구장에는 투구 추적 장비가 있습니다", 19, SOFT)], [body("공 하나하나의 구속과 폼까지 볼 수 있습니다", 19, SOFT)]],
         align=PP_ALIGN.CENTER, line=1.25)
    notes(s, "[0:15] 혹사를 피해 프로에 가도 부상 위험은 이어집니다. 프로 구장에는 투구 추적 장비가 있어 공 하나하나의 구속과 폼까지 볼 수 있습니다.")


def s7_usual(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "선수끼리 말고, 평소의 나와 비교")
    picture(s, DECK / "usual_band.png", 0.6, 1.85, 6.3)
    rows = (("7", "%", "선수끼리 비교(구속·체중·나이)로 설명되는 팔꿈치 수술"), ("6.7", "회", "'구속 1mph 하락' 규칙: 부상 없는 투수 100등판당 경보"),
            ("1", "회", "피치시그널 설계: 정상이면 100등판에 한 번"))
    for i, (num, unit, label) in enumerate(rows):
        y = 1.9 + i * 1.32
        text(s, 7.4, y, 1.9, 1.0, [[(num, BRUSH, 44, BLUE if i == 2 else INK), (unit, TITLE, 28, BLUE if i == 2 else INK)]], anchor=MSO_ANCHOR.MIDDLE)
        text(s, 9.35, y + 0.12, 3.5, 0.9, [[body(label, 15, INK, i == 2)]], anchor=MSO_ANCHOR.MIDDLE)
    picture(s, DECK / "box_blue.png", 0.55, 6.05, 12.25, 1.05)
    text(s, 0.95, 6.2, 11.5, 0.8, [[body("KBO 구장에도 추적 장비가 있지만 데이터는 공개되지 않습니다. 모든 공이 공개된 MLB 2021~2026으로 만들고 검증했고, 구단 데이터를 넣으면 그대로 작동합니다", 15, WHITE, True)]],
         anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "[1:05] 부상 전에 구속이 떨어진다는 연구는 많습니다. 그런데 구속, 체중, 나이처럼 선수끼리 비교하는 요인으로는 팔꿈치 수술의 7%만 설명됩니다. "
             "그래서 선수끼리가 아니라 그 선수의 평소와 비교합니다. 평소 범위를 배우고, 그보다 낮은 흐름이 이어질 때만 경보를 울립니다. "
             "경보는 드물어야 믿습니다. 흔히 쓰는 '구속 1마일 하락' 규칙은 부상 없는 투수에게도 100등판에 6.7번 울립니다. 피치시그널은 정상이면 100등판에 한 번만 울리도록 설계했습니다. "
             "KBO 구단도 같은 장비를 쓰지만 데이터는 공개되지 않습니다. 그래서 모든 공이 공개된 MLB 2021년부터 2026년까지로 만들고 검증했고, 구단 데이터를 넣으면 그대로 돌아갑니다.")


def s8_demo_pro(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "시연 ② 프로")
    picture(s, assets["replay"], 0.45, 1.7, 6.6)
    picture(s, assets["bullpen"], 7.05, 1.25, 5.85)
    picture(s, assets["scout"], 7.35, 4.2, 5.5)
    for x, y, t in ((0.75, 1.55, "① 투수 시즌 상세 · 경보 카드"), (7.35, 1.1, "② 팀 불펜 현황판"), (7.65, 4.05, "③ KBO 외국인 투수 영입 전 점검")):
        picture(s, DECK / "box_small.png", x, y, 3.3, 0.5)
        text(s, x + 0.2, y + 0.06, 3.0, 0.4, [[body(t, 13, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "[1:20 시연] (시즌 상세) 2025년 Michael Soroka입니다. 첫 여덟 등판으로 평소를 배웁니다. "
             "(경보) 6월 말부터 구속이 예상 범위 아래로 내려가고, 7월 4일 경보가 울립니다. 카드는 구속이 얼마나 낮은지 보여 주고, 점검을 시작하라고만 합니다. 판단은 사람이 합니다. "
             "(IL) 8월 5일 부상자 명단에 올랐습니다. 첫 경보보다 한 달 뒤입니다. "
             "(불펜) 팀 불펜 현황판에서는 연투와 최근 7일 투구를 신호등으로 봅니다. "
             "(영입 점검) 같은 엔진을 KBO 외국인 투수 영입에도 씁니다. 2024년부터 3년간 외국인 투수 94명 중 68명은 미국 기록이 있어서, 영입 전에 경보와 부상 이력을 한 장으로 봅니다.")


def s9_validation(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "계획을 먼저 고정하고, 한 번만 계산했다")
    picture(s, DECK / "box_wide.png", 0.55, 1.85, 6.3, 2.35)
    text(s, 1.0, 2.05, 5.6, 1.2, [[("10.3%", BRUSH, 54, BLUE), ("  vs  ", BODY, 22, WHITE, True), ("3.3%", BRUSH, 40, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.0, 3.3, 5.6, 0.75, [[body("같은 오경보율(100등판당 1회)에서 부상 사례 vs 부상 없는 투수의 경보율 · MLB 2024~2025", 13, WHITE)]])
    text(s, 0.65, 4.4, 1.6, 0.8, [[("2.2", BRUSH, 34, INK), ("배", TITLE, 22, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.2, 4.48, 4.6, 0.7, [[body("경보 뒤 30일 안 팔 부상 IL (봉인 2026에서도 2.0배)", 14, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.65, 5.35, 6.2, 0.45, [[("한계", TITLE, 22, INK)]])
    lims = ("구속이 떨어지지 않는 부상은 못 잡습니다 (107건 중 96건)", "봉인한 2026 시즌에서 주 가설은 재현되지 않았습니다", "진단이 아니라 점검을 시작하는 신호입니다")
    text(s, 0.65, 5.85, 6.3, 1.3, [[body("·  " + t, 14, INK)] for t in lims], line=1.25)
    picture(s, assets["curve"], 7.2, 1.75, 5.7)
    notes(s, "[0:35] 평가 계획은 결과를 보기 전에 공개해 고정했고, 한 번만 계산했습니다. 같은 오경보율에서 부상 사례는 10.3%, 부상 없는 투수는 3.3%가 경보를 받았습니다. "
             "한계도 있습니다. 구속이 떨어지지 않는 부상은 잡지 못하고, 봉인해 둔 2026 시즌에서는 주 가설이 재현되지 않았습니다. "
             "그래서 진단이 아니라 점검을 일찍 시작하게 돕는 신호입니다.")


def s10_next(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, "다음 무대, 2027 경기도 전국체전")
    rows = (("01", "2026.11 ~ 2027.3", "경기도 고교·대학 2~3곳 시즌 파일럿"), ("02", "2027.10.8 ~ 14  ·  주 개최지 화성", "전국체전 운영본부와 각 팀이 같은 현황판"),
            ("03", "2028 ~", "구단 추적 데이터 연동"))
    for i, (n, when, what) in enumerate(rows):
        y = 2.0 + i * 1.45
        circle(s, 0.65, y, 0.95, n, size=24)
        text(s, 1.9, y + 0.02, 5.6, 0.45, [[body(when, 15, BLUE, True)]])
        text(s, 1.9, y + 0.45, 6.2, 0.6, [[body(what, 18, INK, True)]])
    picture(s, DECK / "box_paper.png", 8.35, 1.6, 4.35, 4.6)
    picture(s, assets["qr"], 9.35, 2.0, 2.35, 2.35)
    text(s, 8.6, 4.5, 3.85, 1.4, [[body("지금 바로 써 보기", 16, INK, True)], [body("robinkim0221.github.io/pitch-signal", 12, MUTED)],
                                  [body("김형준 · UNIST · robin1967@unist.ac.kr", 12, MUTED)]], align=PP_ALIGN.CENTER, line=1.25)
    text(s, 8.35, 6.35, 4.35, 0.8, [[("감사합니다", TITLE, 30, INK)]], align=PP_ALIGN.CENTER)
    picture(s, DECK / "brush_o_ink.png", 6.9, 5.9, 1.2, 1.2)
    notes(s, "[0:20] 첫 목표는 2027년 경기도 전국체전입니다. 운영본부와 각 팀이 이 현황판을 매일 함께 보게 하겠습니다. 그다음은 구단 추적 데이터입니다. "
             "QR로 지금 바로 써 보실 수 있습니다. 감사합니다.")


def backup(prs, t: str, img: Path, w: float, note: str):
    s = prs.slides.add_slide(prs.slide_layouts[6]); background(s); header(s)
    title(s, t, size=32)
    h = fit(img, w)
    picture(s, img, (SW - w) / 2, 1.6 + max(0, (5.6 - h) / 2), w)
    s._element.set("show", "0")
    notes(s, note)


def build(out: Path) -> Path:
    FRAMED.mkdir(parents=True, exist_ok=True)
    assets = {
        "card": framed("card", SCREENS / "04b_ball_ratio.png", angle=-2),
        "read": framed("read_grid", READ_GRID, private=True),
        "hs": framed("highschool", SCREENS / "04_highschool.png", angle=1.2, crop_h=900),
        "replay": framed("replay", SCREENS / "01_replay.png", angle=-1, crop_h=960),
        "bullpen": framed("bullpen", SCREENS / "07_bullpen.png", angle=1.5, crop_h=760),
        "scout": framed("scout", SCREENS / "05_scout.png", angle=-1.2, crop_h=600),
        "curve": framed("curve", FIG / "03_opcurve_val.png", angle=1),
        "qr": qr(DECK / "qr.png"),
    }
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    for f in (s1_cover, s2_problem, s3_users, s4_reading, s5_demo_hs, s6_bridge, s7_usual, s8_demo_pro, s9_validation, s10_next):
        f(prs, assets)
    backup(prs, "백업 · 경보 카드", framed("alert", SCREENS / "02_alert_card.png"), 10.5, "(라이브 시연이 막힐 때) 경보 카드를 캡처로 설명합니다.")
    backup(prs, "백업 · 검증 요약 화면", framed("perf", SCREENS / "03_performance.png"), 9.5, "(라이브 시연이 막힐 때) 검증 요약을 캡처로 설명합니다.")
    backup(prs, "백업 · 2025 전국체전 등판 달력", framed("tour", FIG / "06b_tournament.png"), 10.5, "(질문 대비) 규정 위반 0건인 대회에서도 부하가 몇 명에게 몰립니다.")
    out = Path(out); out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def finalize(pptx: Path, pdf: Path | None, png_dir: Path | None = None) -> None:
    """PowerPoint로 글꼴을 포함해 다시 저장하고, 원하면 PDF와 쪽 그림(점검용)을 만든다."""
    p = Path(pptx).resolve()
    extra = f"$p.SaveAs('{Path(pdf).resolve()}', 32); " if pdf else ""
    if png_dir:
        Path(png_dir).mkdir(parents=True, exist_ok=True)
        extra += f"foreach ($s in $p.Slides) {{ $s.Export('{Path(png_dir).resolve()}\\s' + $s.SlideIndex.ToString('00') + '.png', 'PNG', 1600, 900) }}; "
    tmp = p.with_name(p.stem + "_embed.pptx")            # 같은 이름으로 다시 저장하면 글꼴이 포함되지 않아(2026-10-22 확인) 다른 이름으로 저장 뒤 바꾼다
    script = (f"$pp = New-Object -ComObject PowerPoint.Application; $p = $pp.Presentations.Open('{p}', -1, 0, 0); "
              f"$p.SaveAs('{tmp}', 24, -1); {extra}$p.Close(); $pp.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    tmp.replace(p)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--pdf", action="store_true", help="같은 이름의 PDF도 (PowerPoint 필요)")
    ap.add_argument("--png", default=None, help="쪽 그림을 저장할 폴더 (점검용)")
    args = ap.parse_args()
    out = build(Path(args.out))
    finalize(out, out.with_suffix(".pdf") if args.pdf else None, Path(args.png) if args.png else None)
    print("저장:", out)


if __name__ == "__main__":
    main()
